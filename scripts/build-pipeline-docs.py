"""Writes the two pipeline workbooks from the JSON scripts/build-pipeline-docs.mts emits.

Split in two on purpose. The TypeScript half has to run inside the project — it
reads the catalogue loaders, the card refs and the live query derivation, none
of which can be reimplemented here honestly. The .xlsx writing wants a
spreadsheet library, and the obvious npm one (SheetJS's `xlsx`) is deprecated
there and carries known advisories; a docs generator is not worth a
supply-chain risk. openpyxl is standard and already available.

Not run directly. Use:  npm run docs:pipelines
"""

import json
import os
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    sys.exit("openpyxl is required:  python -m pip install openpyxl")

DATA = os.path.join("docs", ".pipeline-data.json")
OUT = {"pokemon": "pipeline-pokemon.xlsx", "onePiece": "pipeline-one-piece.xlsx"}

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=13)


def looks_like_header(row):
    """A short all-text first row is a column header, not prose."""
    return len(row) > 1 and all(isinstance(c, str) and 0 < len(c) <= 40 for c in row)


def write_sheet(ws, rows):
    for r_i, row in enumerate(rows, start=1):
        for c_i, value in enumerate(row, start=1):
            cell = ws.cell(row=r_i, column=c_i, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if r_i == 1 and len(rows) > 1 and looks_like_header(row):
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
            elif r_i == 1:
                cell.font = TITLE_FONT

    widths = {}
    for row in rows:
        for c_i, value in enumerate(row, start=1):
            widths[c_i] = min(72, max(widths.get(c_i, 12), len(str(value or "")) + 2))
    for c_i, width in widths.items():
        ws.column_dimensions[ws.cell(row=1, column=c_i).column_letter].width = width

    ws.freeze_panes = "A2"


def main():
    if not os.path.exists(DATA):
        sys.exit(f"{DATA} not found — run: npm run docs:pipelines")
    with open(DATA, encoding="utf-8") as fh:
        payload = json.load(fh)

    for key, filename in OUT.items():
        wb = Workbook()
        wb.remove(wb.active)
        for sheet in payload[key]:
            ws = wb.create_sheet(sheet["name"][:31])
            write_sheet(ws, sheet["rows"])
        path = os.path.join("docs", filename)
        wb.save(path)
        names = ", ".join(s["name"] for s in payload[key])
        print(f"[docs] {filename}: {len(payload[key])} sheets — {names}")


if __name__ == "__main__":
    main()
